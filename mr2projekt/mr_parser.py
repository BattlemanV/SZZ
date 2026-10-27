from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from .docx_utils import grid_cells, normalize_space, parse_number, table_grid_cols
from .vocw import Norms


CODE_RE = re.compile(r"^\d{4}$")
MODE_ORDER = ("max", "annual", "daily")
MODE_TITLES = {
    "max": "максимальных концентраций (МР)",
    "annual": "среднегодовых концентраций (СГ)",
    "daily": "среднесуточных концентраций (СС)",
}
# Вид расчёта по строке «Расчет: «…»» в шапке отчёта
CALC_KINDS = (
    ("среднесуточн", "daily"),
    ("средних", "annual"),
    ("рассеивания", "max"),
)
REPORT_TITLE = "УПРЗА «ЭКОЛОГ»"
RESULTS_HEADER = "Концентр. (д. ПДК)"
RESULT_TITLE = "Результаты расчета по веществам (расчетные точки)"
SOURCES_TITLE = "Параметры источников выбросов"
PDK_TITLE = "Расчет проводился по веществам"
GROUPS_TITLE = "Выбросы источников по группам суммации"
EMISSIONS_TITLE = "Выбросы источников по веществам"
# заголовки разделов отчёта: на них заканчивается таблица предыдущего раздела
SECTION_TITLES = (
    "Параметры источников выбросов",
    "Выбросы источников по веществам",
    "Выбросы источников по группам суммации",
    "Расчет проводился по веществам",
    "Результаты расчета по веществам",
    "УПРЗА «ЭКОЛОГ»",
)
GROUP_ROW_RE = re.compile(r"^Группа суммации:\s*(\d{4})")


@dataclass
class Source:
    code: str
    kind: str
    name: str
    height: str = ""
    diameter: str = ""
    volume: str = ""
    speed: str = ""
    temperature: str = ""

    @property
    def is_organized(self) -> bool:
        """Организованные источники нумеруются с 0, неорганизованные — с 6."""
        if self.code.startswith("6"):
            return False
        if self.code.startswith("0"):
            return True
        return self.kind == "1"

    @property
    def title(self) -> str:
        """Наименование без номера в квадратных скобках: «[1] дефлектор» -> «дефлектор»."""
        return re.sub(r"^\s*\[[^\]]*\]\s*", "", self.name).strip()


@dataclass
class ResultPoint:
    fraction: float | None
    mg_m3: float | None
    number: int = 0


@dataclass
class Substance:
    code: str
    name: str
    g_s: float = 0.0
    t_year: float = 0.0
    hazard: str = ""
    max_type: str = "-"
    max_value: str = "-"
    annual_type: str = "-"
    annual_value: str = "-"
    daily_type: str = "-"
    daily_value: str = "-"
    background: bool = False  # «Фоновая концентр. — Учет: Да» в МР

    @property
    def has_annual(self) -> bool:
        return self.annual_type not in {"", "-"}

    @property
    def has_max(self) -> bool:
        """Критерий — максимально-разовая ПДК (ОБУВ не считается)."""
        return self.max_type == "ПДК м/р" and self.max_value not in {"", "-"}

    @property
    def has_daily(self) -> bool:
        if self.daily_type not in {"", "-"}:
            return True
        return self.annual_type == "ПДК с/с"

    @property
    def daily_effective_type(self) -> str:
        if self.daily_type not in {"", "-"}:
            return self.daily_type
        if self.annual_type == "ПДК с/с":
            return self.annual_type
        return "-"

    @property
    def daily_effective_value(self) -> str:
        if self.daily_type not in {"", "-"}:
            return self.daily_value
        if self.annual_type == "ПДК с/с":
            return self.annual_value
        return "-"


@dataclass
class Stats:
    total_g: float = 0.0
    total_t: float = 0.0
    annual_count: int = 0
    annual_g: float = 0.0
    annual_t: float = 0.0
    max_count: int = 0
    max_g: float = 0.0
    max_t: float = 0.0
    daily_count: int = 0
    daily_g: float = 0.0
    daily_t: float = 0.0
    obuv_count: int = 0
    obuv_g: float = 0.0
    obuv_t: float = 0.0


@dataclass
class MRModel:
    path: Path
    sources: list[Source] = field(default_factory=list)
    substances: list[Substance] = field(default_factory=list)
    groups: dict[str, str] = field(default_factory=dict)
    group_totals: dict[str, float] = field(default_factory=dict)
    group_members: dict[str, list[str]] = field(default_factory=dict)
    results: dict[str, dict[str, list[ResultPoint]]] = field(default_factory=dict)
    points_count: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def by_code(self) -> dict[str, Substance]:
        return {item.code: item for item in self.substances}

    @property
    def organized_count(self) -> int:
        return sum(1 for source in self.sources if source.is_organized)

    @property
    def unorganized_count(self) -> int:
        return sum(1 for source in self.sources if not source.is_organized)

    def compute_stats(self) -> Stats:
        stats = Stats(
            total_g=sum(item.g_s for item in self.substances),
            total_t=sum(item.t_year for item in self.substances),
        )
        annual = [item for item in self.substances if item.has_annual]
        maximum = [item for item in self.substances if item.has_max]
        daily = [item for item in self.substances if item.has_daily]
        obuv = [item for item in self.substances if item.max_type == "ОБУВ"]

        stats.annual_count = len(annual)
        stats.annual_g = sum(item.g_s for item in annual)
        stats.annual_t = sum(item.t_year for item in annual)
        stats.max_count = len(maximum)
        stats.max_g = sum(item.g_s for item in maximum)
        stats.max_t = sum(item.t_year for item in maximum)
        stats.daily_count = len(daily)
        stats.daily_g = sum(item.g_s for item in daily)
        stats.daily_t = sum(item.t_year for item in daily)
        stats.obuv_count = len(obuv)
        stats.obuv_g = sum(item.g_s for item in obuv)
        stats.obuv_t = sum(item.t_year for item in obuv)
        return stats


def _section_tables(document, title: str) -> list:
    """Таблица раздела и её продолжения (до заголовка следующего раздела)."""
    tables = document.tables
    start = next(
        (
            index
            for index, table in enumerate(tables)
            if _first_cell_text(table).startswith(title)
        ),
        None,
    )
    if start is None:
        return []
    result = [tables[start]]
    for table in tables[start + 1 :]:
        if _first_cell_text(table).startswith(SECTION_TITLES):
            break
        result.append(table)
    return result


def _parse_sources_and_emissions(tables) -> tuple[list[Source], list[Substance]]:
    """Источники и выбросы веществ (таблица может продолжаться в следующих)."""
    sources: dict[str, Source] = {}
    substances: dict[str, Substance] = {}
    order: list[str] = []

    rows = [
        grid_cells(row, table_grid_cols(table))
        for table in tables
        for row in table.rows
    ]
    for values in rows:
        if len(values) < 9:
            continue
        code = (values[0] or "").strip()
        if not CODE_RE.match(code):
            continue

        sign = (values[1] or "").strip()
        if sign in {"+", "-", "%"}:
            if code not in sources:
                sources[code] = Source(
                    code=code,
                    kind=(values[3] or "").strip(),
                    name=normalize_space(values[4]),
                    height=(values[6] or "").strip(),
                    diameter=(values[7] or "").strip(),
                    volume=(values[10] or "").strip(),
                    speed=(values[13] or "").strip(),
                    temperature=(values[15] or "").strip(),
                )
            continue

        g_s = parse_number(values[5])
        t_year = parse_number(values[8])
        if g_s is None or t_year is None:
            continue
        name = normalize_space(values[1])
        if not name:
            continue
        if code not in substances:
            substances[code] = Substance(code=code, name=name)
            order.append(code)
        substances[code].g_s += g_s
        substances[code].t_year += t_year

    return list(sources.values()), [substances[code] for code in order]


def _parse_emission_totals(document) -> dict[str, float]:
    """«Итого» г/с по веществам из раздела «Выбросы источников по веществам»."""
    totals: dict[str, float] = {}
    current: str | None = None
    for table in _section_tables(document, EMISSIONS_TITLE):
        ncols = table_grid_cols(table)
        for row in table.rows:
            values = [normalize_space(item) for item in grid_cells(row, ncols)]
            match = re.match(r"Вещество:\s*(\d{4})", values[0])
            if match:
                current = match.group(1)
                continue
            if current and values[0].startswith("Итого"):
                number = next(
                    (
                        parse_number(item)
                        for item in values[1:]
                        if parse_number(item) is not None
                    ),
                    None,
                )
                if number is not None:
                    totals[current] = number
                current = None
    return totals


def _check_emissions(substances: list[Substance], totals: dict[str, float]) -> list[str]:
    """Сверка суммы выбросов по источникам с «Итого» «Эколога»."""
    warnings: list[str] = []
    by_code = {item.code: item for item in substances}
    for code, total in totals.items():
        item = by_code.get(code)
        if item is None:
            warnings.append(
                f"Вещество {code} есть в «Выбросах источников по веществам», "
                "но не найдено в параметрах источников"
            )
        elif abs(item.g_s - total) > max(1e-5, 1e-4 * total):
            warnings.append(
                f"Вещество {code}: сумма выбросов по источникам {item.g_s:.7f} г/с, "
                f"а «Итого» в отчёте {total:.7f} г/с"
            )
    return warnings


def _background_column(table, ncols: int) -> int | None:
    """Столбец «Учет» (фоновая концентрация) в шапке таблицы ПДК."""
    for row in table.rows[:4]:
        for index, text in enumerate(grid_cells(row, ncols)):
            if normalize_space(text) == "Учет":
                return index
    return None


def _parse_pdk(table) -> tuple[dict[str, Substance], dict[str, str], set[str]]:
    ncols = table_grid_cols(table)
    pdk: dict[str, Substance] = {}
    groups: dict[str, str] = {}
    background: set[str] = set()
    column = _background_column(table, ncols)

    for row in table.rows[3:]:
        values = grid_cells(row, ncols)
        if len(values) < 8:
            continue
        code = (values[0] or "").strip()
        if not CODE_RE.match(code):
            continue
        name = normalize_space(values[1])
        max_type = normalize_space(values[2]) or "-"
        annual_type = normalize_space(values[4]) or "-"
        daily_type = normalize_space(values[6]) or "-"
        if column is not None and normalize_space(values[column]).lower() == "да":
            background.add(code)
        if "группа" in max_type.lower():
            groups[code] = name
            continue
        pdk[code] = Substance(
            code=code,
            name=name,
            max_type=max_type,
            max_value=normalize_space(values[3]) or "-",
            annual_type=annual_type,
            annual_value=normalize_space(values[5]) or "-",
            daily_type=daily_type,
            daily_value=normalize_space(values[7]) or "-",
            background=code in background,
        )
    return pdk, groups, background


def _first_cell_text(table) -> str:
    if not table.rows:
        return ""
    return normalize_space(table.rows[0].cells[0].text)


def _find_table(document, title: str, fallback_index: int):
    for table in document.tables:
        if _first_cell_text(table).startswith(title):
            return table
    return document.tables[fallback_index]


@dataclass
class ReportBlock:
    """Один отчёт «Эколога» внутри файла (сводный файл содержит несколько)."""

    kind: str | None
    calc: str = ""
    enterprise: str = ""
    tables: list = field(default_factory=list)


def _calc_kind(calc: str) -> str | None:
    lower = calc.lower()
    return next((kind for marker, kind in CALC_KINDS if marker in lower), None)


def _report_blocks(document) -> list[ReportBlock]:
    blocks: list[ReportBlock] = []
    for table in document.tables:
        if _first_cell_text(table).startswith(REPORT_TITLE):
            lines = [normalize_space(row.cells[0].text) for row in table.rows]
            calc = next((line for line in lines if line.startswith("Расчет:")), "")
            if calc or not blocks:
                enterprise = next(
                    (line for line in lines if line.startswith("Предприятие:")), ""
                )
                blocks.append(ReportBlock(_calc_kind(calc), calc, enterprise))
            continue
        if not blocks:
            blocks.append(ReportBlock(None))
        blocks[-1].tables.append(table)
    return blocks


def _is_results_table(table) -> bool:
    if _first_cell_text(table).startswith(RESULT_TITLE):
        return True
    return any(
        normalize_space(cell.text).startswith(RESULTS_HEADER)
        for row in table.rows[:6]
        for cell in row.cells
    )


def _parse_result_tables(tables) -> dict[str, list[ResultPoint]]:
    """Результаты по расчётным точкам одного отчёта."""
    start = next(
        (index for index, table in enumerate(tables) if _is_results_table(table)),
        None,
    )
    sequences: dict[str, list[ResultPoint]] = {}
    if start is None:
        return sequences

    current: list[ResultPoint] | None = None
    for table in tables[start:]:
        ncols = table_grid_cols(table)
        if ncols < 5:
            continue
        for row in table.rows:
            values = grid_cells(row, ncols)
            if not values:
                continue
            first = (values[0] or "").strip()
            if first.startswith("Вещество:"):
                match = re.search(r"(\d{4})", first)
                current = (
                    sequences.setdefault(match.group(1), []) if match else None
                )
                continue
            if current is None or not first.isdigit():
                continue
            point_no = int(first)
            if not 1 <= point_no <= 100:
                continue
            fraction = parse_number(values[4]) if len(values) > 4 else None
            mg_m3 = parse_number(values[5]) if len(values) > 5 else None
            current.append(ResultPoint(fraction=fraction, mg_m3=mg_m3, number=point_no))

    # В отчётах СС и СГ точки идут не по порядку номеров
    for sequence in sequences.values():
        sequence.sort(key=lambda point: point.number)
    return sequences


def _collect_results(
    max_document,
    files: dict[str, Path | None],
) -> tuple[dict[str, dict[str, list[ResultPoint]]], list[str]]:
    """Результаты МР/СГ/СС: из отдельных файлов, иначе — из сводного МР."""
    warnings: list[str] = []
    main_blocks = _report_blocks(max_document)
    enterprise = main_blocks[0].enterprise if main_blocks else ""
    if main_blocks and main_blocks[0].kind not in {"max", None}:
        warnings.append(
            "Файл МР: первым идёт не расчёт максимальных концентраций "
            f"({main_blocks[0].calc})"
        )

    # Сначала файл, указанный для режима, затем остальные файлы, затем МР
    sources = [
        (mode, path, _report_blocks(Document(str(path))))
        for mode, path in files.items()
        if path is not None
    ]
    results: dict[str, dict[str, list[ResultPoint]]] = {}
    for mode in MODE_ORDER:
        candidates = sorted(sources, key=lambda item: item[0] != mode)
        candidates.append((None, None, main_blocks))
        block = None
        for field_mode, path, blocks in candidates:
            block = next((item for item in blocks if item.kind == mode), None)
            if block is not None:
                source = f"файл {path.name}" if path else "файл МР"
                if field_mode not in {None, mode}:
                    warnings.append(
                        f"{source} указан не в том поле: в нём расчёт "
                        f"{MODE_TITLES[mode]}, использован для него"
                    )
                break
        if block is None and mode == "max" and main_blocks:
            block, source = main_blocks[0], "файл МР"
        if block is None:
            given = files.get(mode)
            where = f"в файле {given.name}" if given else "ни в одном файле"
            warnings.append(f"Нет расчёта {MODE_TITLES[mode]} ({where})")
            results[mode] = {}
            continue
        if enterprise and block.enterprise and block.enterprise != enterprise:
            warnings.append(
                f"{source}: другое предприятие ({block.enterprise}), "
                f"в МР — {enterprise}"
            )
        results[mode] = _parse_result_tables(block.tables)
    return results, warnings


@dataclass
class ReportFiles:
    mr: Path
    annual: Path | None = None
    daily: Path | None = None
    skipped: list[Path] = field(default_factory=list)


def classify_reports(paths: list[str | Path]) -> ReportFiles:
    """Разложить файлы по видам расчёта (МР/СГ/СС) по шапкам отчётов.

    Сводный файл, который начинается с МР, считается отчётом МР.
    Файлы без шапки «УПРЗА «ЭКОЛОГ»» пропускаются.
    """
    found: dict[str, list[Path]] = {mode: [] for mode in MODE_ORDER}
    skipped: list[Path] = []
    for item in paths:
        path = Path(item)
        blocks = [
            block
            for block in _report_blocks(Document(str(path)))
            if block.kind is not None
        ]
        if not blocks:
            skipped.append(path)
            continue
        found[blocks[0].kind].append(path)

    if not found["max"]:
        raise ValueError(
            "Среди файлов нет отчёта МР (расчёт рассеивания, максимальные "
            "концентрации)"
        )
    for mode, files in found.items():
        if len(files) > 1:
            names = ", ".join(path.name for path in files)
            raise ValueError(
                f"Несколько отчётов {MODE_TITLES[mode]}: {names}. "
                "Перетащите по одному отчёту каждого вида."
            )
    return ReportFiles(
        mr=found["max"][0],
        annual=(found["annual"] or [None])[0],
        daily=(found["daily"] or [None])[0],
        skipped=skipped,
    )


def _check_results(model: MRModel) -> list[str]:
    """Сверка состава веществ в расчётах СГ и СС с отчётом МР."""
    warnings: list[str] = []
    codes = {item.code for item in model.substances}
    for mode, attr in (("annual", "annual_type"), ("daily", "daily_type")):
        mode_results = model.results.get(mode) or {}
        if not mode_results:
            continue
        missing = [
            item.code
            for item in model.substances
            if getattr(item, attr) not in {"", "-"} and item.code not in mode_results
        ]
        extra = sorted(code for code in mode_results if code not in codes)
        if missing:
            warnings.append(
                f"В расчёте {MODE_TITLES[mode]} нет веществ: {', '.join(missing)}"
            )
        if extra:
            warnings.append(
                f"В расчёте {MODE_TITLES[mode]} есть вещества, которых нет в МР: "
                f"{', '.join(extra)}"
            )
    return warnings


def fmt_norm(value: float) -> str:
    """Записать норматив так, как его печатает «Эколог»: 0.04, 5E-5, 1E-6."""
    text = f"{value:g}"
    if "e" in text:
        mantissa, exponent = text.split("e")
        text = f"{mantissa}E{int(exponent)}"
    return text


def _same_norm(mr_type: str, mr_value: str, norm: tuple[str, float] | None) -> bool:
    value = parse_number(mr_value)
    if norm is None:
        # «ПДК с/с» со значением «-» Эколог печатает, когда норматива нет
        return mr_type in {"", "-"} or value is None
    return (
        mr_type == norm[0]
        and value is not None
        and abs(value - norm[1]) <= 1e-9 * max(abs(norm[1]), 1.0)
    )


def _apply_norms(substances: list[Substance], norms: dict[str, Norms]) -> list[str]:
    """Нормативы и класс опасности берутся из справочника VocW."""
    warnings: list[str] = []
    for substance in substances:
        info = norms.get(substance.code)
        if info is None:
            warnings.append(
                f"VocW: вещество {substance.code} не найдено, нормативы взяты из MR"
            )
            continue
        if info.unit != "мг/м3":
            warnings.append(
                f"VocW: норматив {substance.code} задан в «{info.unit}», "
                "в мг/м3 не пересчитан"
            )

        modes = (
            ("max", "максимальные", info.max_norm,
             substance.max_type, substance.max_value),
            ("annual", "среднегодовые", info.annual_norm,
             substance.annual_type, substance.annual_value),
            ("daily", "среднесуточные", info.daily_norm,
             substance.daily_effective_type, substance.daily_effective_value),
        )
        for mode, title, norm, mr_type, mr_value in modes:
            if not _same_norm(mr_type, mr_value, norm):
                vocw_text = "-" if norm is None else f"{norm[0]} {fmt_norm(norm[1])}"
                warnings.append(
                    f"VocW: {substance.code} {title}: в MR {mr_type} {mr_value}, "
                    f"в справочнике {vocw_text}"
                )
            norm_type, norm_value = (
                ("-", "-") if norm is None else (norm[0], fmt_norm(norm[1]))
            )
            setattr(substance, f"{mode}_type", norm_type)
            setattr(substance, f"{mode}_value", norm_value)
        substance.hazard = info.hazard_label
    return warnings


def _parse_group_emissions(
    document,
) -> tuple[dict[str, list[str]], dict[str, float]]:
    """Состав и «Итого» групп суммации, как их посчитал «Эколог»."""
    members: dict[str, list[str]] = {}
    totals: dict[str, float] = {}
    tables = document.tables
    start = next(
        (
            index
            for index, table in enumerate(tables)
            if _first_cell_text(table).startswith(GROUPS_TITLE)
        ),
        None,
    )
    if start is None:
        return members, totals

    current: str | None = None
    for table in tables[start:]:
        if _first_cell_text(table).startswith(PDK_TITLE):
            break
        ncols = table_grid_cols(table)
        for row in table.rows:
            values = grid_cells(row, ncols)
            if len(values) < 6:
                continue
            first = normalize_space(values[0])
            match = GROUP_ROW_RE.match(first)
            if match:
                current = match.group(1)
                members.setdefault(current, [])
                continue
            if current is None:
                continue
            if first.startswith("Итого"):
                total = parse_number(values[5])
                if total is not None:
                    totals[current] = total
                continue
            code = (values[4] or "").strip()
            if CODE_RE.match(code) and code not in members[current]:
                members[current].append(code)
    return members, totals


def point_numbers(results: dict[str, dict[str, list[ResultPoint]]]) -> list[int]:
    """Номера расчётных точек по возрастанию (без номера — по позиции)."""
    numbers = {
        point.number or index + 1
        for mode in results.values()
        for points in mode.values()
        for index, point in enumerate(points)
    }
    return sorted(numbers)


def points_by_number(points: list[ResultPoint]) -> dict[int, ResultPoint]:
    by_number: dict[int, ResultPoint] = {}
    for index, point in enumerate(points):
        by_number.setdefault(point.number or index + 1, point)
    return by_number


def _compute_group_totals(
    substances: list[Substance], members: dict[str, list[str]]
) -> dict[str, float]:
    by_code = {item.code: item for item in substances}
    totals: dict[str, float] = {}
    for group_code, member_codes in members.items():
        totals[group_code] = sum(
            by_code[code].g_s for code in member_codes if code in by_code
        )
    return totals


def parse_mr(
    path: str | Path,
    norms: dict[str, Norms] | None = None,
    group_members: dict[str, list[str]] | None = None,
    *,
    annual_path: str | Path | None = None,
    daily_path: str | Path | None = None,
    hazard_classes: dict[str, str] | None = None,
) -> MRModel:
    """Разобрать отчёт МР и (если заданы) отдельные отчёты СГ и СС.

    Источники, выбросы, ПДК и группы суммации берутся из отчёта МР;
    нормативы и классы опасности уточняются по справочнику VocW (norms),
    а hazard_classes (файл пользователя) переопределяет классы опасности.
    """
    path = Path(path)
    document = Document(str(path))
    sources, substances = _parse_sources_and_emissions(
        _section_tables(document, SOURCES_TITLE)
    )
    pdk, groups, background = _parse_pdk(_find_table(document, PDK_TITLE, 9))

    for substance in substances:
        info = pdk.get(substance.code)
        if info is None:
            continue
        substance.name = info.name or substance.name
        substance.max_type = info.max_type
        substance.max_value = info.max_value
        substance.annual_type = info.annual_type
        substance.annual_value = info.annual_value
        substance.daily_type = info.daily_type
        substance.daily_value = info.daily_value
        substance.background = info.background

    pdk_order = {code: index for index, code in enumerate(pdk)}
    substances.sort(
        key=lambda item: pdk_order.get(item.code, len(pdk_order))
    )

    results, result_warnings = _collect_results(
        document,
        {
            "annual": Path(annual_path) if annual_path else None,
            "daily": Path(daily_path) if daily_path else None,
        },
    )
    if not results["max"]:
        raise ValueError("В отчёте МР не найдены результаты расчёта по точкам")
    points_count = len(point_numbers(results))
    model = MRModel(
        path=path,
        sources=sources,
        substances=substances,
        groups=groups,
        results=results,
        points_count=points_count,
        warnings=result_warnings,
    )
    model.warnings.extend(_check_emissions(substances, _parse_emission_totals(document)))
    if norms:
        model.warnings.extend(_apply_norms(substances, norms))
    for substance in substances:
        if hazard_classes and substance.code in hazard_classes:
            substance.hazard = hazard_classes[substance.code]
    model.warnings.extend(_check_results(model))

    mr_members, mr_totals = _parse_group_emissions(document)
    vocw_members = group_members or {}
    emitted = {item.code for item in substances}
    computed = _compute_group_totals(substances, vocw_members)
    for group_code in groups:
        reference = vocw_members.get(group_code)
        if reference is None and vocw_members:
            model.warnings.append(f"VocW: нет состава группы суммации {group_code}")
        members = mr_members.get(group_code)
        if members is not None and reference is not None:
            expected = sorted(code for code in reference if code in emitted)
            if sorted(members) != expected:
                model.warnings.append(
                    f"Группа {group_code}: в MR {', '.join(sorted(members))}, "
                    f"по VocW {', '.join(expected) or '-'}; сумма взята из MR"
                )
        model.group_members[group_code] = list(members or reference or [])
        if group_code in mr_totals:
            model.group_totals[group_code] = mr_totals[group_code]
        elif group_code in computed:
            model.group_totals[group_code] = computed[group_code]
    return model
