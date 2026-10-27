from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from .docx_utils import (
    fmt_emission,
    fmt_ru,
    normalize_space,
    replace_regex_in_paragraph,
    replace_regex_with_suffix,
    set_cell_text,
    set_paragraph_text,
)
from .mr_parser import MRModel, Stats
from .table52 import _strip_trailing_parens, rebuild as rebuild_table52


def _table31_templates(table) -> tuple[object | None, object | None]:
    substance_tr = None
    group_tr = None
    for row in table.rows[3:]:
        if len(row._tr.tc_lst) >= 11 and substance_tr is None:
            substance_tr = copy.deepcopy(row._tr)
        elif len(row._tr.tc_lst) < 11 and group_tr is None:
            group_tr = copy.deepcopy(row._tr)
    return substance_tr, group_tr


def _clear_rows(table, start: int) -> None:
    for row in list(table.rows[start:]):
        table._tbl.remove(row._tr)


def _append_row(table, template_tr):
    tr = copy.deepcopy(template_tr)
    table._tbl.append(tr)
    return table.rows[-1]


def _code_column(table) -> int:
    """Номер столбца «Код» в шапке таблицы 3.1 (1 — если слева есть «№ п/п»)."""
    for index, cell in enumerate(table.rows[0].cells):
        if normalize_space(cell.text) == "Код":
            return index
    return 0


def _fill_table_31(document, model: MRModel) -> list[str]:
    warnings: list[str] = []
    table = document.tables[0]
    substance_tr, group_tr = _table31_templates(table)

    if model.substances and substance_tr is None:
        return ["3.1: в шаблоне нет строки вещества"]
    if model.groups and group_tr is None:
        return ["3.1: в шаблоне нет строки группы суммы"]

    _clear_rows(table, 3)

    no_hazard = [item.code for item in model.substances if not item.hazard]
    if no_hazard:
        warnings.append(
            f"3.1: нет класса опасности для веществ {', '.join(no_hazard)} — "
            "впишите вручную или добавьте в файл «классы опасности.csv» "
            "в папке программы"
        )

    # Столбец «№ п/п» (если есть в шаблоне) стоит слева от «Код»
    offset = _code_column(table)
    number = 0

    def put(cells, index: int, text: str) -> None:
        set_cell_text(cells[offset + index], text)

    for substance in model.substances:
        row = _append_row(table, substance_tr)
        cells = row.cells
        number += 1
        if offset:
            set_cell_text(cells[0], str(number))
        max_type = substance.max_type if substance.max_type not in {"", "-"} else "ПДК м/р"
        annual_type = (
            substance.annual_type if substance.annual_type not in {"", "-"} else "ПДК с/г"
        )
        put(cells, 0, substance.code)
        put(cells, 1, normalize_space(substance.name))
        put(cells, 2, substance.hazard)
        put(cells, 3, max_type)
        put(cells, 4, substance.max_value)
        put(cells, 5, annual_type)
        put(cells, 6, substance.annual_value)
        put(cells, 7, "ПДК с/с")
        put(cells, 8, substance.daily_effective_value)
        put(cells, 9, fmt_emission(substance.g_s, GS_DECIMALS))
        put(cells, 10, fmt_emission(substance.t_year, T_DECIMALS))
        if len(cells) > offset + 11:
            put(cells, 11, "")

    for group_code, group_name in model.groups.items():
        row = _append_row(table, group_tr)
        if offset:
            # группы суммации не нумеруются: «№ п/п» объединён с кодом группы
            row.cells[0].merge(row.cells[offset])
        cells = row.cells
        total = model.group_totals.get(group_code)
        if total is None:
            warnings.append(f"3.1: нет состава группы {group_code}")
        put(cells, 0, group_code)
        put(cells, 1, normalize_space(group_name))
        put(cells, 2, "" if total is None else fmt_emission(total, GS_DECIMALS))
        if len(cells) > offset + 11:
            put(cells, 11, "")

    return warnings


# Точность выбросов как в справочнике «Эколога» (Setting.ADT: DigitsGS, DigitsTG)
GS_DECIMALS = 7
T_DECIMALS = 6


PLACEHOLDER_RE = re.compile(r"…[….]*|\.{3,}[….]*")


def _find_paragraph(paragraphs, marker: str):
    marker = marker.lower()
    return next((item for item in paragraphs if marker in item.text.lower()), None)


def _fill_placeholders(paragraph, values: list[str], warnings: list[str]) -> bool:
    """Заменить многоточия в абзаце значениями по порядку."""
    text = paragraph.text
    count = len(PLACEHOLDER_RE.findall(text))
    if not count:
        return False
    if count != len(values):
        warnings.append(
            f"Текст: в абзаце «{normalize_space(text)[:40]}…» "
            f"{count} пропусков, ожидалось {len(values)}"
        )
    iterator = iter(values)
    set_paragraph_text(
        paragraph,
        PLACEHOLDER_RE.sub(lambda match: next(iterator, match.group(0)), text),
    )
    return True


# абзац об ОБУВ в шаблоне ищется по любой из этих фраз
OBUV_MARKERS = ("критерием оценки является ОБУВ", "(ОБУВ)")
OBUV_TEXT = (
    "Для {count} {noun}, критерием оценки является ОБУВ и объем суммарного "
    "выброса составляет {grams} г\\с и {tons} т\\год."
)

# «для 1 веществ» -> «для 1 вещества» (и прилагательные перед словом)
SINGULAR_FORMS = (
    (" выделяющихся загрязняющих веществ", " выделяющегося загрязняющего вещества"),
    (" загрязняющих веществ", " загрязняющего вещества"),
    (" веществ", " вещества"),
)
COUNT_RE = re.compile(
    r"(?<![\d,])(\d+)("
    + "|".join(re.escape(plural) for plural, _single in SINGULAR_FORMS)
    + r")\b"
)


def _agree_substances(text: str) -> str:
    """Согласовать «N веществ» с числом: 1, 21, 31… — единственное число."""

    def repl(match: re.Match) -> str:
        count, phrase = int(match.group(1)), match.group(2)
        if count % 10 == 1 and count % 100 != 11:
            phrase = dict(SINGULAR_FORMS)[phrase]
        return match.group(1) + phrase

    return COUNT_RE.sub(repl, text)


def _fix_agreement(paragraph) -> None:
    text = paragraph.text
    fixed = _agree_substances(text)
    if fixed != text:
        set_paragraph_text(paragraph, fixed)


def _update_obuv(paragraphs, stats: Stats, warnings: list[str]) -> None:
    """Абзац о веществах с ОБУВ — после абзаца о среднегодовой ПДК."""
    existing = next(
        (
            item
            for marker in OBUV_MARKERS
            if (item := _find_paragraph(paragraphs, marker)) is not None
        ),
        None,
    )
    if stats.obuv_count == 0:
        if existing is not None:
            existing._p.getparent().remove(existing._p)
        return

    values = [
        str(stats.obuv_count),
        fmt_ru(stats.obuv_g, 5),
        fmt_ru(stats.obuv_t, 5),
    ]
    if existing is not None:
        if not _fill_placeholders(existing, values, warnings):
            replace_regex_with_suffix(existing, r"(\d+)(\s+веществ)", values[:1])
            replace_regex_in_paragraph(existing, r"(\d+,\d+)", values[1:])
        _fix_agreement(existing)
        return

    anchor = _find_paragraph(paragraphs, "среднегодовая ПДК")
    if anchor is None:
        warnings.append("Текст: некуда вставить абзац о веществах с ОБУВ")
        return
    element = copy.deepcopy(anchor._p)
    anchor._p.addnext(element)
    set_paragraph_text(
        Paragraph(element, anchor._parent),
        OBUV_TEXT.format(
            count=values[0],
            noun=_plural(stats.obuv_count, "вещества", "веществ", "веществ"),
            grams=values[1],
            tons=values[2],
        ),
    )


SUBSTANCES_MARKER = "выделяются"
SOURCES_MARKER = "Согласно проектным данным"


def _plural(count: int, one: str, few: str, many: str) -> str:
    """Форма слова по числу: 1 источник, 2 источника, 5 источников."""
    if count % 10 == 1 and count % 100 != 11:
        return one
    if 2 <= count % 10 <= 4 and not 12 <= count % 100 <= 14:
        return few
    return many


def _lower_first(text: str) -> str:
    # «Бензин» -> «бензин», но аббревиатуры («ПАУ») не трогаем
    if len(text) > 1 and text[1].isupper():
        return text
    return text[:1].lower() + text[1:]


def _substances_list(model: MRModel, intro_names: dict[str, str]) -> str:
    names: list[str] = []
    for substance in model.substances:
        name = intro_names.get(substance.code) or _lower_first(
            _strip_trailing_parens(normalize_space(substance.name))
        )
        if name and name not in names:
            names.append(name)
    return ", ".join(names)


def _sources_part(sources, kind: str, suffix: str = "") -> str:
    """«4 неорганизованных источника выбросов – проезд, склад»."""
    count = len(sources)
    text = f"{count} " + _plural(
        count, f"{kind}ый источник", f"{kind}ых источника", f"{kind}ых источников"
    ) + suffix
    titles: list[str] = []
    for source in sources:
        title = _lower_first(normalize_space(source.title))
        if title and title not in titles:
            titles.append(title)
    if titles:
        text += " – " + ", ".join(_collapse_titles(titles))
    return text


NUMBERED_TITLE_RE = re.compile(r"^(.*?)(\d+)$")
HUSHING = "гкхжчшщ"


def _plural_word(word: str) -> str:
    """Множественное число первого слова: дефлектор -> дефлекторы, труба -> трубы."""
    lower = word.lower()
    if (
        word != lower  # аббревиатура (ДК, ДГУ)
        or not re.fullmatch(r"[а-яё]+", lower)
        or lower.endswith(("ой", "ый", "ий", "ая", "яя"))  # прилагательное
    ):
        return word
    if lower.endswith("а"):
        return word[:-1] + ("и" if lower[-2:-1] in HUSHING else "ы")
    if lower.endswith(("я", "ь", "й")):
        return word[:-1] + "и"
    if lower[-1] in "бвгджзклмнпрстфхцчшщ":
        return word + ("и" if lower[-1] in HUSHING else "ы")
    return word


def _number_ranges(numbers: list[tuple[int, str]]) -> list[tuple[str, str]]:
    """[(1,'1'),(2,'2'),(3,'3'),(5,'5')] -> [('1','3'), ('5','5')]."""
    ranges: list[tuple[str, str]] = []
    start = previous = numbers[0]
    for item in numbers[1:]:
        if item[0] == previous[0] + 1:
            previous = item
            continue
        ranges.append((start[1], previous[1]))
        start = previous = item
    ranges.append((start[1], previous[1]))
    return ranges


def _is_numbering(prefix: str) -> bool:
    """Число в конце — номер, а не количество.

    «дефлектор ВЕ1», «ДК1», «труба №1», «свеча 2» — номера; а в «гаражные
    боксы 6» число — скорее количество, такие названия не сворачиваем.
    """
    head, _, label = prefix.rpartition(" ")
    if label:  # метка слитно с номером: «ВЕ1», «№1»
        return True
    first = head.split(" ")[0]
    return _plural_word(first) != first  # существительное в ед. числе


def _collapse_titles(titles: list[str]) -> list[str]:
    """«дефлектор ВЕ1» … «дефлектор ВЕ14» -> «дефлекторы ВЕ1–ВЕ14»."""
    groups: dict[str, list[tuple[int, str]]] = {}
    for title in titles:
        match = NUMBERED_TITLE_RE.match(title)
        if match and match.group(1).strip():
            groups.setdefault(match.group(1), []).append(
                (int(match.group(2)), match.group(2))
            )

    result: list[str] = []
    done: set[str] = set()
    for title in titles:
        match = NUMBERED_TITLE_RE.match(title)
        prefix = match.group(1) if match else None
        if (
            prefix is None
            or len(groups.get(prefix, [])) < 2
            or not _is_numbering(prefix)
        ):
            result.append(title)
            continue
        if prefix in done:
            continue
        done.add(prefix)
        # «дефлектор ВЕ» -> слова «дефлекторы» и метка номера «ВЕ»;
        # «ДК» (номер без пробела) -> только метка: «ДК1–ДК3»
        head, _, label = prefix.rpartition(" ")
        parts = ", ".join(
            f"{label}{first}" if first == last else f"{label}{first}–{label}{last}"
            for first, last in _number_ranges(sorted(groups[prefix]))
        )
        if head:
            words = head.split(" ")
            words[0] = _plural_word(words[0])
            result.append(f"{' '.join(words)} {parts}")
        else:
            result.append(parts)
    return result


def _sources_sentence(model: MRModel) -> str:
    """«Согласно проектным данным на предприятии выявлен 1 неорганизованный …»."""
    unorganized = [item for item in model.sources if not item.is_organized]
    organized = [item for item in model.sources if item.is_organized]

    parts: list[tuple[int, str]] = []
    if unorganized:
        parts.append(
            (len(unorganized), _sources_part(unorganized, "неорганизованн", " выбросов"))
        )
    if organized:
        parts.append((len(organized), _sources_part(organized, "организованн")))
    if not parts:
        return f"{SOURCES_MARKER} на предприятии источники выбросов не выявлены."
    verb = _plural(parts[0][0], "выявлен", "выявлено", "выявлено")
    return (
        f"{SOURCES_MARKER} на предприятии {verb} "
        + " и ".join(part for _count, part in parts)
        + "."
    )


def _update_intro(paragraphs, model: MRModel, intro_names: dict[str, str]) -> None:
    """Раздел 3.1: перечень выделяющихся веществ и источники выбросов."""
    substances = _find_paragraph(paragraphs, SUBSTANCES_MARKER)
    if substances is not None:
        text = substances.text
        head = text[: text.lower().index(SUBSTANCES_MARKER) + len(SUBSTANCES_MARKER)]
        if text[len(head):].lstrip().startswith(":"):
            head += ":"  # двоеточие после «выделяются» — как в шаблоне
        set_paragraph_text(
            substances, f"{head} {_substances_list(model, intro_names)}."
        )

    sources = _find_paragraph(paragraphs, SOURCES_MARKER)
    if sources is not None:
        set_paragraph_text(sources, _sources_sentence(model))


def _update_narrative(
    document, model: MRModel, stats: Stats, intro_names: dict[str, str]
) -> list[str]:
    warnings: list[str] = []
    paragraphs = document.paragraphs
    _update_intro(paragraphs, model, intro_names)

    sources = _find_paragraph(paragraphs, "инвентаризации источников")
    if sources is not None and not _fill_placeholders(
        sources,
        [str(model.organized_count), str(model.unorganized_count)],
        warnings,
    ):
        replace_regex_with_suffix(
            sources, r"(\d+)(\s+организованных)", [str(model.organized_count)]
        )
        replace_regex_with_suffix(
            sources, r"(\d+)(\s+неорганизованных)", [str(model.unorganized_count)]
        )

    total = _find_paragraph(paragraphs, "Общее количество загрязняющих веществ")
    if total is not None and not _fill_placeholders(
        total, [fmt_ru(stats.total_t, 5)], warnings
    ):
        replace_regex_in_paragraph(total, r"(\d+,\d+)", [fmt_ru(stats.total_t, 5)])

    maximum = _find_paragraph(paragraphs, "Как видно из таблицы 3.1")
    max_values = [
        str(len(model.substances)),
        str(stats.max_count),
        fmt_ru(stats.max_g, 5),
        fmt_ru(stats.max_t, 5),
    ]
    if maximum is not None and len(PLACEHOLDER_RE.findall(maximum.text)) == 3:
        # короткий вариант: «для … выделяющихся загрязняющих веществ» без общего числа
        max_values = max_values[1:]
    if maximum is not None and not _fill_placeholders(
        maximum, max_values, warnings
    ):
        replace_regex_with_suffix(
            maximum,
            r"(\d+)(\s+веществ)",
            [str(len(model.substances)), str(stats.max_count)],
        )
        replace_regex_in_paragraph(
            maximum,
            r"(\d+,\d+)",
            [fmt_ru(stats.max_g, 5), fmt_ru(stats.max_t, 5)],
        )

    for marker, count, grams, tons in (
        ("среднесуточная ПДК", stats.daily_count, stats.daily_g, stats.daily_t),
        ("среднегодовая ПДК", stats.annual_count, stats.annual_g, stats.annual_t),
    ):
        paragraph = _find_paragraph(paragraphs, marker)
        if paragraph is None:
            continue
        values = [str(count), fmt_ru(grams, 5), fmt_ru(tons, 5)]
        if not _fill_placeholders(paragraph, values, warnings):
            replace_regex_with_suffix(paragraph, r"(\d+)(\s+веществ)", values[:1])
            replace_regex_in_paragraph(paragraph, r"(\d+,\d+)", values[1:])
        _fix_agreement(paragraph)
    if maximum is not None:
        _fix_agreement(maximum)

    for marker in (
        "Общее количество загрязняющих веществ",
        "Как видно из таблицы 3.1",
        "среднесуточная ПДК",
        "среднегодовая ПДК",
    ):
        if _find_paragraph(paragraphs, marker) is None:
            warnings.append(f"Текст: не найден абзац «{marker}»")

    _update_obuv(paragraphs, stats, warnings)
    return warnings


def generate(
    model: MRModel,
    template_path: str | Path,
    output_path: str | Path,
    labels: dict[str, str] | None = None,
    intro_names: dict[str, str] | None = None,
) -> dict[str, object]:
    document = Document(str(template_path))
    stats = model.compute_stats()
    warnings = list(model.warnings)
    warnings.extend(_fill_table_31(document, model))
    warnings.extend(rebuild_table52(document, model, labels))
    warnings.extend(_update_narrative(document, model, stats, intro_names or {}))
    document.save(str(output_path))
    return {
        "output": str(output_path),
        "stats": stats.__dict__,
        "organized": model.organized_count,
        "unorganized": model.unorganized_count,
        "points": model.points_count,
        "warnings": warnings,
    }
