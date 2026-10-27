from __future__ import annotations

import re
from typing import Iterator

from docx.oxml.ns import qn
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


def cell_text(tc) -> str:
    parts: list[str] = []
    for node in tc.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag in {qn("w:br"), qn("w:cr"), qn("w:tab")}:
            parts.append(" ")
        elif node.tag == qn("w:p") and parts and not parts[-1].endswith(" "):
            parts.append(" ")
    return "".join(parts)


def table_grid_cols(table: Table) -> int:
    return len(table._tbl.tblGrid.gridCol_lst)


def iter_grid_cells(row) -> Iterator[tuple[int, int, _Cell]]:
    """Yield (grid_start, span, cell) for real cells, including merged ones."""
    pos = 0
    for tc in row._tr.tc_lst:
        span = 1
        if tc.tcPr is not None and tc.tcPr.gridSpan is not None:
            span = int(tc.tcPr.gridSpan.val)
        yield pos, span, _Cell(tc, row)
        pos += span


def grid_cells(row, ncols: int) -> list[str]:
    values = [""] * ncols
    for start, span, cell in iter_grid_cells(row):
        text = cell_text(cell._tc)
        for index in range(start, min(start + span, ncols)):
            values[index] = text
    return values


def cells_by_grid_start(row) -> dict[int, _Cell]:
    return {start: cell for start, _span, cell in iter_grid_cells(row)}


def set_cell_text(cell: _Cell, text: str) -> None:
    """Set cell text while keeping the first run's formatting."""
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)

    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace paragraph text, keeping the first run's font."""
    font = None
    size = None
    if paragraph.runs:
        font = paragraph.runs[0].font.name
        size = paragraph.runs[0].font.size
    paragraph.text = text
    for run in paragraph.runs:
        if font:
            run.font.name = font
        if size:
            run.font.size = size


def replace_in_paragraph(paragraph: Paragraph, old: str, new: str) -> bool:
    text = paragraph.text
    if old not in text:
        return False
    set_paragraph_text(paragraph, text.replace(old, new))
    return True


def replace_regex_in_paragraph(paragraph: Paragraph, pattern: str, replacements) -> None:
    text = paragraph.text
    iterator = iter(replacements)

    def repl(_match):
        try:
            return next(iterator)
        except StopIteration:
            return _match.group(0)

    new_text = re.sub(pattern, repl, text)
    if new_text != text:
        set_paragraph_text(paragraph, new_text)


def replace_regex_with_suffix(paragraph: Paragraph, pattern: str, replacements) -> None:
    """Replace group 1 with the next value and keep group 2 as suffix."""
    text = paragraph.text
    iterator = iter(replacements)

    def repl(match):
        try:
            value = next(iterator)
        except StopIteration:
            return match.group(0)
        suffix = match.group(2) if match.lastindex and match.lastindex >= 2 else ""
        return f"{value}{suffix}"

    new_text = re.sub(pattern, repl, text)
    if new_text != text:
        set_paragraph_text(paragraph, new_text)


def parse_number(value: str | None) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if text in {"", "-", "\u2014", "\u2013"}:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def fmt_emission(value: float, decimals: int) -> str:
    """Выброс как в «Экологе» (г/с — 7 знаков, т/год — 6); меньше 1E-6 — 2E-7."""
    if value == 0:
        return "0"
    if abs(value) < 1e-6:
        mantissa, exponent = f"{value:.6e}".split("e")
        mantissa = mantissa.rstrip("0").rstrip(".")
        return f"{mantissa}E{int(exponent)}"
    text = f"{value:.{decimals}f}".rstrip("0").rstrip(".")
    return text or "0"


def fmt_ru(value: float, digits: int = 5) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def fmt_result(value: float | None) -> str:
    if value is None:
        return "-"
    if abs(value) < 0.01:
        return f"{value:.2E}"
    return f"{value:.2f}"
