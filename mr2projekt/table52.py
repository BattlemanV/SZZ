from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from typing import Iterable

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .docx_utils import (
    cell_text,
    fmt_result,
    iter_grid_cells,
    normalize_space,
    set_cell_text,
    set_paragraph_text,
)
from .mr_parser import MRModel, point_numbers, points_by_number

MODE_LIMIT = 10
BLOCK_LIMIT = 5
FONT_SIZE = Pt(10)
NUMBER_WIDTH = 295
SUBSTANCE_MODE_WIDTH = 879
GROUP_MODE_WIDTH = 1701
AVAILABLE_WIDTH = 9354

CODE_RE = re.compile(r"^\s*(\d{3,4})(?!\d)")


@dataclass(frozen=True)
class ModeColumn:
    mode: str
    label: str


@dataclass(frozen=True)
class Block:
    code: str
    label: str
    modes: tuple[ModeColumn, ...]
    is_group: bool = False


@dataclass(frozen=True)
class ParsedBlock:
    code: str
    label: str
    is_group: bool
    modes: dict[str, int]


@dataclass(frozen=True)
class ParsedResult:
    block: ParsedBlock
    values: dict[int, dict[str, str]] = field(default_factory=dict)


def _short_code(code: str) -> str:
    return str(int(code)) if code.isdigit() else code


def _strip_trailing_parens(text: str) -> str:
    """Убрать пояснения в скобках в конце названия (с вложенными скобками)."""
    text = text.strip()
    while text.endswith(")"):
        depth = 0
        for index in range(len(text) - 1, -1, -1):
            if text[index] == ")":
                depth += 1
            elif text[index] == "(":
                depth -= 1
                if depth == 0:
                    break
        if depth != 0 or index == 0:
            break
        text = text[:index].rstrip()
    return text


BACKGROUND_SUFFIX = " с учетом фона"


def _block_label(
    code: str,
    name: str,
    labels: dict[str, str],
    is_group: bool,
    background: bool = False,
) -> str:
    if is_group:
        label = f"{_short_code(code)} Группа суммации"
    elif code in labels:
        label = labels[code]
    else:
        text = _strip_trailing_parens(normalize_space(name))
        label = f"{_short_code(code)} {text}".strip()
    # в МР для вещества учитывался фон («Фоновая концентр. — Учет: Да»)
    return label + BACKGROUND_SUFFIX if background else label


def _has_fractions(model: MRModel, mode: str, code: str) -> bool:
    """В MR есть результаты в долях ПДК (не прочерки) для режима."""
    points = model.results.get(mode, {}).get(code) or []
    return any(point.fraction is not None for point in points)


def build_blocks(
    model: MRModel, labels: dict[str, str] | None = None
) -> tuple[list[Block], list[Block], list[str]]:
    labels = labels or {}
    warnings: list[str] = []
    substances: list[Block] = []
    groups: list[Block] = []

    for substance in model.substances:
        # Колонка режима нужна, только если у вещества есть норматив для
        # этого режима и «Эколог» выдал по нему доли ПДК.
        candidates = (
            ("max", substance.max_type),
            ("daily", substance.daily_type),
            ("annual", substance.annual_type),
        )
        modes = [
            ModeColumn(mode, norm_type)
            for mode, norm_type in candidates
            if norm_type not in {"", "-"}
            and _has_fractions(model, mode, substance.code)
        ]
        if not modes:
            warnings.append(f"5.2: у вещества {substance.code} нет режимов")
            continue
        substances.append(
            Block(
                code=substance.code,
                label=_block_label(
                    substance.code,
                    substance.name,
                    labels,
                    is_group=False,
                    background=substance.background,
                ),
                modes=tuple(modes),
            )
        )

    for group_code, group_name in model.groups.items():
        if not _has_fractions(model, "max", group_code):
            warnings.append(f"5.2: нет результатов для группы {group_code}")
            continue
        groups.append(
            Block(
                code=group_code,
                # у групп суммации «с учетом фона» не пишем
                label=_block_label(group_code, group_name, labels, is_group=True),
                modes=(ModeColumn("max", ""),),
                is_group=True,
            )
        )

    return substances, groups, warnings


def pack_blocks(
    blocks: Iterable[Block],
    mode_limit: int = MODE_LIMIT,
    block_limit: int = BLOCK_LIMIT,
) -> list[list[Block]]:
    items = list(blocks)
    memo: dict[int, tuple[list[list[Block]], int]] = {}

    def solve(index: int) -> tuple[list[list[Block]], int]:
        if index >= len(items):
            return [], 0
        if index in memo:
            return memo[index]

        best: tuple[list[list[Block]], int] | None = None
        mode_count = 0
        for end in range(index, min(len(items), index + block_limit)):
            mode_count += len(items[end].modes)
            if mode_count > mode_limit and end > index:
                break
            rest, rest_score = solve(end + 1)
            candidate = ([items[index : end + 1]] + rest, mode_count**2 + rest_score)
            if best is None or (
                len(candidate[0]),
                candidate[1],
            ) < (
                len(best[0]),
                best[1],
            ):
                best = candidate
        assert best is not None
        memo[index] = best
        return best

    return solve(0)[0]


def _code_from_label(text: str) -> str | None:
    match = CODE_RE.match(text)
    if match is None:
        return None
    return match.group(1).zfill(4)


def find_result_tables(document) -> list[Table]:
    tables: list[Table] = []
    for table in document.tables:
        if len(table.rows) < 3:
            continue
        if table.rows[0].cells[0].text.strip() != "№":
            continue
        if not any(
            start > 0 and _code_from_label(cell_text(cell._tc)) is not None
            for start, _span, cell in iter_grid_cells(table.rows[1])
        ):
            continue
        tables.append(table)
    return tables


def _parse_mode_cells(table: Table, start: int, end: int) -> list[tuple[int, str]]:
    cells = []
    for cell_start, _span, cell in iter_grid_cells(table.rows[2]):
        if start <= cell_start < end:
            cells.append((cell_start, normalize_space(cell_text(cell._tc))))
    return cells


def _modes_for_block(
    code: str,
    label: str,
    mode_cells: list[tuple[int, str]],
) -> ParsedBlock:
    is_group = "групп" in label.lower() or "суммации" in label.lower()
    modes: dict[str, int] = {}

    if is_group:
        if mode_cells:
            modes["max"] = mode_cells[0][0]
        return ParsedBlock(code, label, True, modes)

    max_start: int | None = None
    non_max: list[tuple[int, str]] = []
    for cell_start, mode_label in mode_cells:
        lower = mode_label.lower()
        if "м/р" in lower or "обув" in lower:
            max_start = cell_start
        else:
            non_max.append((cell_start, mode_label))

    if max_start is not None:
        modes["max"] = max_start
    if len(non_max) == 1:
        modes["annual"] = non_max[0][0]
    elif len(non_max) >= 2:
        daily = next(
            (start for start, text in non_max if "с/с" in text.lower()),
            non_max[0][0],
        )
        annual = next(
            (start for start, text in non_max if "с/г" in text.lower()),
            non_max[-1][0],
        )
        if daily == annual:
            daily, annual = non_max[0][0], non_max[1][0]
        modes["daily"] = daily
        modes["annual"] = annual

    if not modes and mode_cells:
        modes["annual"] = mode_cells[0][0]
    return ParsedBlock(code, label, False, modes)


def parse_results(document) -> list[ParsedResult]:
    parsed: list[ParsedResult] = []
    for table in find_result_tables(document):
        blocks: list[ParsedBlock] = []
        for start, span, cell in iter_grid_cells(table.rows[1]):
            if start == 0:
                continue
            label = normalize_space(cell_text(cell._tc))
            code = _code_from_label(label)
            if code is None:
                continue
            blocks.append(
                _modes_for_block(
                    code,
                    label,
                    _parse_mode_cells(table, start, start + span),
                )
            )

        rows: dict[int, dict[int, str]] = {}
        for row in table.rows[3:]:
            first = next(iter_grid_cells(row), None)
            if first is None:
                continue
            point_text = normalize_space(cell_text(first[2]._tc))
            if not point_text.isdigit():
                continue
            point = int(point_text)
            rows[point] = {
                start: normalize_space(cell_text(cell._tc))
                for start, _span, cell in iter_grid_cells(row)
            }

        for block in blocks:
            values: dict[int, dict[str, str]] = {}
            for point, cells in rows.items():
                values[point] = {
                    mode: cells.get(start, "") for mode, start in block.modes.items()
                }
            parsed.append(ParsedResult(block, values))
    return parsed


def _prepare_cell(tc, width: int, span: int) -> None:
    tcPr = tc.get_or_add_tcPr()
    for tag in ("w:vMerge", "w:hMerge"):
        element = tcPr.find(qn(tag))
        if element is not None:
            tcPr.remove(element)

    tc_w = tcPr.get_or_add_tcW()
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")

    if span > 1:
        grid_span = tcPr.get_or_add_gridSpan()
        grid_span.set(qn("w:val"), str(span))
    elif tcPr.gridSpan is not None:
        tcPr._remove_gridSpan()

    for proof in list(tc.iter(qn("w:proofErr"))):
        proof.getparent().remove(proof)


def _prepare_row_template(tr):
    tr = copy.deepcopy(tr)
    if tr.trPr is not None:
        for tag in ("w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter"):
            element = tr.trPr.find(qn(tag))
            if element is not None:
                tr.trPr.remove(element)
    return tr


BORDER_SIDES = ("top", "left", "bottom", "right")


def _set_border(tc, side: str, visible: bool = True) -> None:
    """Задать границу ячейки: одинарная линия либо её отсутствие."""
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    element = borders.find(qn(f"w:{side}"))
    if element is None:
        element = OxmlElement(f"w:{side}")
        # порядок дочерних элементов tcBorders: top, left, bottom, right
        index = BORDER_SIDES.index(side)
        following = [
            borders.find(qn(f"w:{name}")) for name in BORDER_SIDES[index + 1 :]
        ]
        anchor = next((item for item in following if item is not None), None)
        if anchor is None:
            borders.append(element)
        else:
            anchor.addprevious(element)
    for key in list(element.attrib):
        del element.attrib[key]
    if visible:
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")
    else:
        element.set(qn("w:val"), "nil")


def _merge_number_header(rows) -> None:
    """Объединить «№» в трёх строках заголовка, текст — по центру."""
    for index, row in enumerate(rows):
        tc = row._tr.tc_lst[0]
        tc.get_or_add_tcPr().vMerge_val = "restart" if index == 0 else "continue"
        if index > 0:
            _set_border(tc, "top", False)
        if index < len(rows) - 1:
            _set_border(tc, "bottom", False)
        cell = _Cell(tc, row)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _merge_group_headers(table: Table, blocks: list[Block]) -> None:
    """У группы суммации норматива нет: название на две строки заголовка."""
    label_cells = table.rows[1]._tr.tc_lst[1:]
    mode_cells = table.rows[2]._tr.tc_lst[1:]
    position = 0
    for block, label_tc in zip(blocks, label_cells):
        span = len(block.modes)
        if block.is_group and span == 1:
            mode_tc = mode_cells[position]
            label_tc.get_or_add_tcPr().vMerge_val = "restart"
            mode_tc.get_or_add_tcPr().vMerge_val = "continue"
            _set_border(label_tc, "bottom", False)
            _set_border(mode_tc, "top", False)
        position += span


def _apply_font_size(table: Table, size) -> None:
    """Один размер шрифта для всей таблицы, включая пустые ячейки."""
    for tc in table._tbl.iter(qn("w:tc")):
        for p in tc.iter(qn("w:p")):
            paragraph = Paragraph(p, None)
            for run in paragraph.runs:
                run.font.size = size
            # знак абзаца задаёт высоту строки в пустых ячейках
            pPr = p.get_or_add_pPr()
            rPr = pPr.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                pPr.append(rPr)
            rPr.sz_val = size


def _append_row(table: Table, tbl, template_tr, specs, texts) -> None:
    tr = _prepare_row_template(template_tr)
    for tc in list(tr.tc_lst):
        tr.remove(tc)
    for template_tc, width, span in specs:
        tc = copy.deepcopy(template_tc)
        _prepare_cell(tc, width, span)
        tr.append(tc)
    tbl.append(tr)
    row = table.rows[-1]
    for (_start, _span, cell), text in zip(iter_grid_cells(row), texts):
        set_cell_text(cell, text)


def _set_grid(tbl, widths: list[int]) -> None:
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")


def _table_widths(blocks: list[Block]) -> list[int]:
    group = bool(blocks) and blocks[0].is_group
    mode_count = sum(len(block.modes) for block in blocks)
    if group:
        mode_width = GROUP_MODE_WIDTH
    else:
        available = AVAILABLE_WIDTH - NUMBER_WIDTH
        mode_width = max(
            SUBSTANCE_MODE_WIDTH,
            available // mode_count if mode_count else SUBSTANCE_MODE_WIDTH,
        )
    return [NUMBER_WIDTH] + [mode_width for _ in range(mode_count)]


def _build_table(
    document,
    base_tbl,
    templates,
    cell_templates,
    blocks: list[Block],
    model: MRModel,
    numbers: list[int],
):
    row0, row1, row2, row3 = templates
    num_header, header, label, mode, number, data = cell_templates
    widths = _table_widths(blocks)
    mode_width = widths[1] if len(widths) > 1 else SUBSTANCE_MODE_WIDTH
    mode_count = len(widths) - 1

    tbl = copy.deepcopy(base_tbl)
    for tr in list(tbl.tr_lst):
        tbl.remove(tr)
    _set_grid(tbl, widths)
    table = Table(tbl, document)

    _append_row(
        table,
        tbl,
        row0,
        [(num_header, widths[0], 1), (header, sum(widths[1:]), mode_count)],
        ["№", "Концентрации (д. ПДК)"],
    )

    label_specs = [(number, widths[0], 1)]
    label_texts = [""]
    for block in blocks:
        label_specs.append((label, mode_width * len(block.modes), len(block.modes)))
        label_texts.append(block.label)
    _append_row(table, tbl, row1, label_specs, label_texts)

    mode_specs = [(number, widths[0], 1)]
    mode_texts = [""]
    for block in blocks:
        for mode_column in block.modes:
            mode_specs.append((mode, mode_width, 1))
            mode_texts.append(mode_column.label)
    _append_row(table, tbl, row2, mode_specs, mode_texts)

    # граница между наименованием вещества и применяемым нормативом
    for tc in table.rows[1]._tr.tc_lst[1:]:
        _set_border(tc, "bottom")
    for tc in table.rows[2]._tr.tc_lst[1:]:
        _set_border(tc, "top")
    _merge_number_header(table.rows[:3])
    _merge_group_headers(table, blocks)

    lookups = {
        (block.code, mode_column.mode): points_by_number(
            model.results.get(mode_column.mode, {}).get(block.code) or []
        )
        for block in blocks
        for mode_column in block.modes
    }
    for point_number in numbers:
        specs = [(number, widths[0], 1)]
        texts = [str(point_number)]
        bold = [False]
        for block in blocks:
            for mode_column in block.modes:
                specs.append((data, mode_width, 1))
                point = lookups[(block.code, mode_column.mode)].get(point_number)
                fraction = None if point is None else point.fraction
                texts.append(fmt_result(fraction))
                # превышение ПДК (больше 1) — жирным
                bold.append(fraction is not None and fraction > 1)
        _append_row(table, tbl, row3, specs, texts)
        for cell, is_bold in zip(table.rows[-1].cells, bold):
            for run in cell.paragraphs[0].runs:
                run.bold = True if is_bold else None

    _apply_font_size(table, FONT_SIZE)
    return tbl


def _paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _remove_old_tables(document, old_tables: list[Table]) -> None:
    for table in old_tables:
        parent = table._tbl.getparent()
        if parent is None:
            continue
        previous = table._tbl.getprevious()
        while previous is not None and previous.tag == qn("w:p"):
            text = normalize_space(_paragraph_text(previous))
            if not text:
                before = previous.getprevious()
                previous.getparent().remove(previous)
                previous = before
                continue
            if "продолжение" in text.lower() and "5.2" in text:
                previous.getparent().remove(previous)
            break
        parent.remove(table._tbl)


def rebuild(
    document,
    model: MRModel,
    labels: dict[str, str] | None = None,
) -> list[str]:
    warnings: list[str] = []
    old_tables = find_result_tables(document)
    if not old_tables:
        return ["5.2: не найдены таблицы 5.2"]

    title = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if normalize_space(paragraph.text).lower().startswith("таблица 5.2")
        ),
        None,
    )
    if title is None:
        return ["5.2: не найдена подпись таблицы 5.2"]

    continuation = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if "продолжение" in paragraph.text.lower()
            and "5.2" in paragraph.text
        ),
        None,
    )

    base_table = old_tables[-1]
    if len(base_table.rows) < 4:
        return ["5.2: в шаблоне нет строки-образца"]

    templates = tuple(
        copy.deepcopy(base_table._tbl.tr_lst[index]) for index in range(4)
    )
    cell_templates = (
        copy.deepcopy(base_table.rows[0]._tr.tc_lst[0]),
        copy.deepcopy(base_table.rows[0]._tr.tc_lst[1]),
        copy.deepcopy(base_table.rows[1]._tr.tc_lst[1]),
        copy.deepcopy(base_table.rows[2]._tr.tc_lst[1]),
        copy.deepcopy(base_table.rows[3]._tr.tc_lst[0]),
        copy.deepcopy(base_table.rows[3]._tr.tc_lst[1]),
    )
    base_tbl = copy.deepcopy(base_table._tbl)
    continuation_template = copy.deepcopy(
        (continuation or title)._p
    )
    blank_template = next(
        (
            copy.deepcopy(paragraph._p)
            for paragraph in document.paragraphs
            if not normalize_space(paragraph.text)
            and paragraph._p.find(".//" + qn("w:sectPr")) is None
        ),
        None,
    )
    if blank_template is None:
        blank_template = copy.deepcopy(title._p)

    substances, groups, build_warnings = build_blocks(model, labels)
    warnings.extend(build_warnings)

    anchor = title._p
    following = anchor.getnext()
    if following is not None and following.tag == qn("w:p"):
        anchor = following

    _remove_old_tables(document, old_tables)

    if not substances and not groups:
        warnings.append("5.2: нет веществ и групп")
        return warnings

    numbers = point_numbers(model.results)
    packs = pack_blocks(substances)
    if groups:
        packs.extend(pack_blocks(groups, mode_limit=BLOCK_LIMIT))

    for index, blocks in enumerate(packs):
        if index > 0:
            blank = Paragraph(copy.deepcopy(blank_template), title._parent)
            set_paragraph_text(blank, "")
            blank.paragraph_format.page_break_before = False
            anchor.addnext(blank._p)
            anchor = blank._p

            paragraph = Paragraph(
                copy.deepcopy(continuation_template), title._parent
            )
            set_paragraph_text(paragraph, "Продолжение Таблицы 5.2.")
            # части таблицы идут одна за другой; подпись не отрывается от таблицы
            paragraph.paragraph_format.page_break_before = False
            paragraph.paragraph_format.keep_with_next = True
            anchor.addnext(paragraph._p)
            anchor = paragraph._p
        tbl = _build_table(
            document,
            base_tbl,
            templates,
            cell_templates,
            blocks,
            model,
            numbers,
        )
        anchor.addnext(tbl)
        anchor = tbl

    if model.points_count <= 0:
        warnings.append("5.2: нет расчётных точек")
    return warnings
