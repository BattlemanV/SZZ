from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from mr2projekt.generator import (
    PLACEHOLDER_RE,
    _agree_substances,
    _collapse_titles,
    _plural,
    _sources_sentence,
    _substances_list,
)
from mr2projekt.pipeline import convert
from mr2projekt.mr_parser import (
    MRModel,
    Source,
    Substance,
    _collect_results,
    _parse_result_tables,
    classify_reports,
    parse_mr,
)
from mr2projekt.reference import load_hazard_classes, load_reference
from mr2projekt.table52 import find_result_tables, parse_results


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "отчет выход.docx"


class StandaloneTests(unittest.TestCase):
    """Проверки, которым не нужны отчёты «Эколога» и справочник."""

    def test_unsorted_points(self) -> None:
        """В отчётах СС и СГ точки идут не по порядку — после разбора отсортированы."""
        document = Document()
        table = document.add_table(rows=0, cols=6)
        table.add_row().cells[0].text = (
            "Результаты расчета по веществам (расчетные точки)"
        )
        table.add_row().cells[0].text = "Вещество: 0123 Железа оксид"
        for point_no in (3, 1, 2):
            cells = table.add_row().cells
            cells[0].text = str(point_no)
            cells[4].text = f"0.{point_no}"
            cells[5].text = "1"

        points = _parse_result_tables(document.tables)["0123"]
        self.assertEqual([point.number for point in points], [1, 2, 3])
        self.assertEqual([point.fraction for point in points], [0.1, 0.2, 0.3])

    def test_result_sections_with_continuations(self) -> None:
        document = Document()

        def header(calc: str) -> None:
            table = document.add_table(rows=0, cols=1)
            table.add_row().cells[0].text = "УПРЗА «ЭКОЛОГ» 4.70"
            table.add_row().cells[0].text = "Предприятие: 1, Тест"
            table.add_row().cells[0].text = f"Расчет: «{calc}»"

        header("Расчет рассеивания по МРР-2017")
        first = document.add_table(rows=0, cols=6)
        first.add_row().cells[0].text = (
            "Результаты расчета по веществам (расчетные точки)"
        )
        first.add_row().cells[0].text = "Вещество: 0123 Железа оксид"
        for point_no in (1, 2):
            cells = first.add_row().cells
            cells[0].text = str(point_no)
            cells[4].text = f"0.{point_no}"
            cells[5].text = str(point_no)

        continuation = document.add_table(rows=0, cols=6)
        for point_no in (3, 4):
            cells = continuation.add_row().cells
            cells[0].text = str(point_no)
            cells[4].text = f"0.{point_no}"
            cells[5].text = str(point_no)

        # СС раньше СГ и без заголовка «Результаты расчета…»: режим — по шапке
        for calc, fraction in (
            ("Расчет среднесуточных концентраций", "0.6"),
            ("Расчет средних концентраций по МРР-2017", "0.5"),
        ):
            header(calc)
            section = document.add_table(rows=0, cols=6)
            section.add_row().cells[0].text = "Вещество: 0123 Железа оксид"
            row = section.add_row().cells
            row[4].text = "Концентр. (д. ПДК)"
            cells = section.add_row().cells
            cells[0].text = "1"
            cells[4].text = fraction
            cells[5].text = "1"

        results, warnings = _collect_results(document, {})

        self.assertEqual(warnings, [])
        self.assertEqual(len(results["max"]["0123"]), 4)
        self.assertEqual(results["annual"]["0123"][0].fraction, 0.5)
        self.assertEqual(results["daily"]["0123"][0].fraction, 0.6)

    def test_classify_reports(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        folder = Path(directory.name)

        def make(name: str, *calcs: str) -> Path:
            document = Document()
            for calc in calcs:
                table = document.add_table(rows=0, cols=1)
                table.add_row().cells[0].text = "УПРЗА «ЭКОЛОГ» 4.70"
                table.add_row().cells[0].text = f"Расчет: «{calc}»"
            document.save(folder / name)
            return folder / name

        mr = make("a.docx", "Расчет рассеивания по МРР-2017")
        sg = make("b.docx", "Расчет средних концентраций по МРР-2017")
        ss = make("c.docx", "Расчет среднесуточных концентраций")
        other = make("d.docx")

        files = classify_reports([ss, other, sg, mr])
        self.assertEqual((files.mr, files.annual, files.daily), (mr, sg, ss))
        self.assertEqual(files.skipped, [other])
        with self.assertRaises(ValueError):
            classify_reports([sg, ss])
        with self.assertRaises(ValueError):
            classify_reports([mr, make("e.docx", "Расчет рассеивания по МРР-2017")])

    def test_hazard_classes_file(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "классы опасности.csv"
        # как сохраняет Excel: cp1251, без ведущих нулей
        path.write_bytes(
            "Код;Класс опасности\n301;3\n2732;нет\n0123;5\nабв;1\n".encode("cp1251")
        )
        self.assertEqual(load_hazard_classes(path), {"0301": "3", "2732": "Нет"})

    def test_reference_without_vocw(self) -> None:
        reference = load_reference(Path(tempfile.gettempdir()) / "нет такого VocW")
        self.assertIsNone(reference.vocw)
        self.assertIn("VocW", reference.vocw_error)
        self.assertEqual(reference.substances, {})
        self.assertEqual(reference.group_members, {})

    def test_plural(self) -> None:
        forms = {n: _plural(n, "источник", "источника", "источников")
                 for n in (1, 2, 5, 11, 12, 21, 22, 25, 111)}
        self.assertEqual(forms, {
            1: "источник", 2: "источника", 5: "источников", 11: "источников",
            12: "источников", 21: "источник", 22: "источника", 25: "источников",
            111: "источников",
        })

    def test_sources_sentence(self) -> None:
        def model(organized: int, names: list[str]) -> MRModel:
            item = MRModel(path=Path("x"))
            item.sources = [
                Source(code=f"{index:04d}", kind="1", name="[1] дефлектор")
                for index in range(1, organized + 1)
            ] + [
                Source(code=f"{6001 + index}", kind="3", name=name)
                for index, name in enumerate(names)
            ]
            return item

        self.assertEqual(
            _sources_sentence(model(0, ["Внутренний проезд и стоянка"])),
            "Согласно проектным данным на предприятии выявлен 1 неорганизованный "
            "источник выбросов – внутренний проезд и стоянка.",
        )
        self.assertEqual(
            _sources_sentence(model(11, ["Проезд", "Склад", "Проезд", "ДК1"])),
            "Согласно проектным данным на предприятии выявлено 4 неорганизованных "
            "источника выбросов – проезд, склад, ДК1 и 11 организованных источников "
            "– дефлектор.",
        )
        self.assertEqual(
            _sources_sentence(model(21, [])),
            "Согласно проектным данным на предприятии выявлен 21 организованный "
            "источник – дефлектор.",
        )

    def test_agreement(self) -> None:
        self.assertEqual(
            _agree_substances("Для 1 веществ, из 21 загрязняющих веществ, для 11 веществ"),
            "Для 1 вещества, из 21 загрязняющего вещества, для 11 веществ",
        )
        self.assertEqual(
            _agree_substances("для 1 выделяющихся загрязняющих веществ"),
            "для 1 выделяющегося загрязняющего вещества",
        )

    def test_collapse_titles(self) -> None:
        cases = [
            (["дефлектор ВЕ1", "дефлектор ВЕ2", "дефлектор ВЕ3", "труба ДГУ"],
             ["дефлекторы ВЕ1–ВЕ3", "труба ДГУ"]),
            (["дефлектор ВЕ1", "дефлектор ВЕ2", "дефлектор ВЕ5", "склад"],
             ["дефлекторы ВЕ1–ВЕ2, ВЕ5", "склад"]),
            (["свеча 1", "свеча 2", "ДК1 (ЛОС)"], ["свечи 1–2", "ДК1 (ЛОС)"]),
            (["ДК1", "ДК2", "ДК3"], ["ДК1–ДК3"]),
            (["дефлектор ВЕ1", "склад"], ["дефлектор ВЕ1", "склад"]),
            # число после слова во множественном числе — количество, не номер
            (["гаражные боксы 6", "гаражные боксы 3"],
             ["гаражные боксы 6", "гаражные боксы 3"]),
        ]
        for titles, expected in cases:
            self.assertEqual(_collapse_titles(titles), expected)

    def test_substances_list(self) -> None:
        item = MRModel(path=Path("x"))
        item.substances = [
            Substance(code="0301", name="Азота диоксид (Двуокись азота)"),
            Substance(code="0304", name="Азот (II) оксид (Азот монооксид)"),
            Substance(code="9999", name="Тестовое вещество (пояснение)"),
        ]
        self.assertEqual(
            _substances_list(item, {"0301": "оксиды азота", "0304": "оксиды азота"}),
            "оксиды азота, тестовое вещество",
        )


FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"


class FixtureTests(unittest.TestCase):
    """Проверки на сохранённых отчётах «Эколога» (tests/fixtures).

    Отчёты — данные заказчиков, в репозиторий они не входят; без них
    эти проверки пропускаются.
    """

    @classmethod
    def setUpClass(cls) -> None:
        if not FIXTURES_DIR.is_dir() or not any(FIXTURES_DIR.glob("*.docx")):
            raise unittest.SkipTest("нет отчётов «Эколога» в tests/fixtures")

    def _convert(self, mr: str, **extra) -> tuple[Document, list[str]]:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        output = Path(directory.name) / "out.docx"
        summary = convert(
            FIXTURES_DIR / mr, TEMPLATE, output,
            with_appendix=False,  # приложение проверяет test_appendix
            **{key: FIXTURES_DIR / value for key, value in extra.items()},
        ).summary
        return Document(str(output)), list(summary["warnings"])

    def test_intro_barrikada(self) -> None:
        document, warnings = self._convert(
            "баррикада_мр.docx",
            annual_path="баррикада_сг.docx",
            daily_path="баррикада_сс.docx",
        )
        texts = [item.text for item in document.paragraphs]
        self.assertEqual(warnings, [])
        self.assertTrue(texts[0].startswith("3.1. Краткая характеристика технологии"))
        self.assertEqual(
            texts[1],
            "При работе предприятия в атмосферный воздух выделяются оксиды азота, "
            "сажа, оксид серы, оксид углерода, бензин, керосин.",
        )
        self.assertEqual(
            texts[3],
            "Согласно проектным данным на предприятии выявлен 1 неорганизованный "
            "источник выбросов – внутренний проезд и стоянка.",
        )
        self.assertIsNone(PLACEHOLDER_RE.search("\n".join(texts)))
        self.assertIn(
            "Для 1 вещества, критерием оценки является ОБУВ и объем суммарного "
            "выброса составляет 0,00428 г\\с и 0,00789 т\\год.",
            texts,
        )
        self.assertTrue(find_result_tables(document))

    def test_table31_numbering(self) -> None:
        document, _warnings = self._convert(
            "баррикада_мр.docx",
            annual_path="баррикада_сг.docx",
            daily_path="баррикада_сс.docx",
        )
        table = document.tables[0]
        self.assertEqual(table.rows[0].cells[0].text, "№ п/п")
        self.assertEqual(table.rows[0].cells[1].text, "Код")
        rows = [(row.cells[0].text, row.cells[1].text) for row in table.rows[3:]]
        self.assertEqual(
            rows,
            [
                ("1", "0301"), ("2", "0304"), ("3", "0328"), ("4", "0330"),
                # у группы «№ п/п» объединён с кодом — одна ячейка на два столбца
                ("5", "0337"), ("6", "2704"), ("7", "2732"), ("6204", "6204"),
            ],
        )
        group_tc = table.rows[-1]._tr.tc_lst[0]
        self.assertEqual(group_tc.grid_span, 2)
        self.assertEqual(table.rows[-1].cells[0].text, "6204")

    def test_background_labels(self) -> None:
        """«Учет» фона в МР -> «с учетом фона» в шапке таблицы 5.2."""
        document, _warnings = self._convert(
            "агзс_мр.docx",
            annual_path="агзс_сг.docx",
            daily_path="агзс_сс.docx",
        )
        labels = {item.block.code: item.block.label for item in parse_results(document)}
        self.assertEqual(labels["0301"], "301 Азота диоксид с учетом фона")
        self.assertEqual(labels["0330"], "330 Сера диоксид с учетом фона")
        self.assertEqual(labels["0337"], "337 Углерода оксид с учетом фона")
        # у групп суммации подпись не нужна, даже если фон учитывался
        self.assertEqual(labels["6204"], "6204 Группа суммации")
        self.assertEqual(labels["0304"], "304 Азота оксид")
        self.assertEqual(labels["2732"], "2732 Керосин")

        # в «Баррикаде» фон нигде не учитывался
        document, _warnings = self._convert("баррикада_мр.docx")
        self.assertFalse(any(
            "фона" in item.block.label for item in parse_results(document)
        ))

    def test_sources_table_continued(self) -> None:
        """Таблица «Параметры источников» продолжается во второй таблице."""
        model = parse_mr(FIXTURES_DIR / "агзс_мр.docx")
        self.assertEqual(len(model.sources), 8)
        self.assertEqual(
            [item.code for item in model.substances],
            ["0301", "0304", "0328", "0330", "0337", "0402", "0415", "0418",
             "2704", "2732"],
        )
        # суммы совпадают с «Итого» «Эколога» — предупреждений о выбросах нет
        self.assertFalse(any("Итого" in item for item in model.warnings))

    def test_appendix(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        output = Path(directory.name) / "Отчет СЗЗ (агзс).docx"
        summary = convert(
            FIXTURES_DIR / "агзс_мр.docx", TEMPLATE, output,
            annual_path=FIXTURES_DIR / "агзс_сг.docx",
            daily_path=FIXTURES_DIR / "агзс_сс.docx",
        ).summary
        self.assertEqual(summary["warnings"], [])
        appendix = Path(summary["appendix"])
        self.assertEqual(appendix.name, "Отчет СЗЗ (агзс) — приложение.docx")

        # основной отчёт — без таблиц «Эколога»
        main = Document(str(output))
        self.assertFalse(any(
            table.rows[0].cells[0].text.startswith("УПРЗА") for table in main.tables
        ))

        document = Document(str(appendix))
        # весь текст приложения — Times New Roman
        fonts = {
            rFonts.get(qn("w:ascii"))
            for root in (document.element.body, document.styles.element)
            for rFonts in root.iter(qn("w:rFonts"))
        }
        self.assertEqual(fonts, {"Times New Roman"})
        tables = document.tables
        heads = [
            table for table in tables
            if table.rows[0].cells[0].text.startswith("УПРЗА")
        ]
        calcs = [
            next(row.cells[0].text for row in table.rows
                 if row.cells[0].text.startswith("Расчет:"))
            for table in heads
        ]
        self.assertEqual(len(calcs), 3)
        self.assertIn("рассеивания", calcs[0])
        self.assertIn("средних", calcs[1])
        self.assertIn("среднесуточных", calcs[2])

        # во всех результатах точки каждого вещества идут по возрастанию
        start = tables.index(heads[0])
        sequences: list[list[int]] = []
        in_results = False
        for table in tables[start:]:
            first = table.rows[0].cells[0].text.strip()
            if first.startswith("УПРЗА"):
                in_results = False  # новый отчёт: до результатов идут другие разделы
            elif first.startswith("Результаты расчета"):
                in_results = True
            if not in_results:
                continue
            for row in table.rows:
                text = row.cells[0].text.strip()
                if text.startswith("Вещество:"):
                    sequences.append([])
                elif text.isdigit() and sequences:
                    sequences[-1].append(int(text))
        point_lists = [items for items in sequences if items]
        self.assertGreater(len(point_lists), 20)
        for items in point_lists:
            self.assertEqual(items, sorted(items))

    def test_no_dash_columns(self) -> None:
        """В 5.2 нет колонок из одних прочерков."""
        document, _warnings = self._convert(
            "агзс_мр.docx",
            annual_path="агзс_сг.docx",
            daily_path="агзс_сс.docx",
        )
        parsed = {item.block.code: item for item in parse_results(document)}
        self.assertEqual(set(parsed["2732"].block.modes), {"max"})  # только ОБУВ
        self.assertEqual(set(parsed["0330"].block.modes), {"max", "annual"})
        for item in parsed.values():
            for mode in item.block.modes:
                column = [values.get(mode, "") for values in item.values.values()]
                self.assertTrue(
                    any(value not in {"", "-"} for value in column),
                    f"{item.block.code}/{mode}: колонка из прочерков",
                )

    def test_table52_formatting(self) -> None:
        document, _warnings = self._convert("жби_мр.docx")
        bold_seen = 0
        for table in find_result_tables(document):
            rows = table._tbl.tr_lst
            merges = [rows[i].tc_lst[0].tcPr.vMerge_val for i in range(3)]
            self.assertEqual(merges, ["restart", "continue", "continue"])

            # граница между наименованием вещества и нормативом
            # (у групп суммации эти ячейки объединены — границы там нет)
            for row_index, side in ((1, "bottom"), (2, "top")):
                for tc in rows[row_index].tc_lst[1:]:
                    if tc.tcPr.vMerge_val is not None:
                        continue
                    border = tc.tcPr.find(qn("w:tcBorders")).find(qn(f"w:{side}"))
                    self.assertEqual(border.get(qn("w:val")), "single")

            for row in table.rows[3:]:
                for cell in row.cells[1:]:
                    runs = cell.paragraphs[0].runs
                    value = float(cell.text) if cell.text not in {"", "-"} else 0
                    self.assertEqual(bool(runs and runs[0].bold), value > 1, cell.text)
                    bold_seen += value > 1
        self.assertGreater(bold_seen, 0)

    def test_intro_zhbi(self) -> None:
        document, warnings = self._convert("жби_мр.docx")
        sources = next(
            item.text for item in document.paragraphs
            if item.text.startswith("Согласно проектным данным")
        )
        self.assertIn("выявлено 15 неорганизованных источников выбросов – ", sources)
        self.assertTrue(
            sources.endswith(
                " и 16 организованных источников – дефлекторы ВЕ1–ВЕ14, труба ДГУ."
            )
        )
        self.assertEqual(sources.count("ДК1 (ЛОС)"), 1)
        # в этом файле только расчёт МР
        self.assertTrue(any("(СГ)" in item for item in warnings))


if __name__ == "__main__":
    unittest.main()
